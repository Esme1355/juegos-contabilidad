"""Genera la planilla de contactos institucionales de clínicas privadas de Santiago
y un CSV compatible con la importación de contactos de Wix."""
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUT = Path(__file__).parent

# Clinica, RazonSocial, RUT, Direccion, Comuna, Telefono, Web,
# EmailGeneral, EmailCompras, CanalCompras(URL/nota), EmailRRHH, CanalRRHH(URL/nota),
# Docencia/Capacitacion interna, NombreContacto, CargoContacto, Fuente
C = [
 ("Clínica Alemana","Clínica Alemana de Santiago S.A.","96.770.100-9","Av. Vitacura 5951","Vitacura","+56 2 2210 1111","clinicaalemana.cl",
  "","","proveedores.alemana.cl (portal facturas Banco de Chile)","","clinicaalemana.trabajando.cl","Área científico-docente; Facultad de Medicina CAS-UDD","","","superdesalud.gob.cl/registro/clinica-alemana-de-santiago/"),
 ("Clínica Dávila","Clínica Dávila y Servicios Médicos SpA","96.530.470-3","Av. Recoleta 464","Recoleta","+56 2 2730 8000","davila.cl",
  "","","davila.cl/proveedores (sin correo, vía teléfono central)","","davila.trabajando.cl","Dirección de Docencia; campus clínico UDD","José Retamal","Director de Docencia","davila.cl/unidad-de-investigacion-clinica/"),
 ("Clínica Dávila Vespucio","Clínica Vespucio S.A. (Grupo Banmédica)","96.898.980-4","Serafín Zamora 190","La Florida","+56 2 2820 6700","davila.cl/davilavespucio",
  "","","davila.cl/proveedores","","davila.trabajando.cl","","","","superdesalud.gob.cl/registro/clinica-davila-vespucio/"),
 ("Clínica Santa María","Clínica Santa María S.A.","90.753.000-0","Av. Santa María 0500","Providencia","+56 2 2913 0000","clinicasantamaria.cl",
  "","","","","clinicasantamaria.trabajando.cl","Centro de Simulación Avanzada","","","superdesalud.gob.cl/registro/clinica-santa-maria/"),
 ("Clínica Bupa Santiago","Clínica Bupa Santiago S.A.","76.242.774-5","Av. Departamental 1455","La Florida","+56 2 3240 5650","clinicabupasantiago.cl",
  "","","bupa.cl/portal-para-proveedores","","clinicabupasantiago.cl/quienes-somos/trabaja-con-nosotros","","","","superdesalud.gob.cl/registro/clinica-bupa-santiago/"),
 ("Clínica Indisa","Instituto de Diagnóstico S.A.","92.051.000-0","Av. Santa María 1810","Providencia","+56 2 2362 5555","indisa.cl",
  "","","indisa.cl/informacion/proveedores","","indisa.trabajando.cl","Investigación y docencia; campo clínico UNAB","","","superdesalud.gob.cl/registro/clinica-indisa/"),
 ("Clínica MEDS","Clínica Meds La Dehesa S.A.","76.336.039-3","Av. José Alcalde Délano 10581","Lo Barnechea","+56 2 2499 6400","meds.cl",
  "contacto@meds.cl","","meds.cl/cadena-de-suministro/","","meds.trabajando.cl","meds.cl/docencia-meds/","","","superdesalud.gob.cl/registro/clinica-meds-la-dehesa/"),
 ("Clínica Universidad de los Andes","Universidad de los Andes","71.614.000-8","Av. Plaza 2501","Las Condes","+56 2 2618 3000","clinicauandes.cl",
  "contacto@clinicauandes.cl","gestionproveedores@clinicauandes.cl","Compras vía SAP Ariba; facturas: atencionproveedores@uandes.cl","","clinicauandes.trabajando.cl","","","","uandes.cl/proveedores/"),
 ("Fundación Arturo López Pérez (FALP)","Fundación Arturo López Pérez","70.377.400-8","José Manuel Infante 805","Providencia","+56 2 2712 8000","falp.org",
  "subgerenciadeventas@falp.org","","proveedores.bancochile.cl (pymes) / globalecbusiness.com","","falp.trabajando.cl","Instituto Oncológico: docencia e investigación","","","falp.org/contactoall-2/"),
 ("Clínica RedSalud Providencia","Empresas RedSalud S.A.","78.040.520-1","Av. Salvador 100","Providencia","+56 2 2366 2055","redsalud.cl",
  "","","","","redsalud.trabajando.cl","","Pilar Torres Medina","Gerente de Personas (corporativo RedSalud)","redsalud.cl/acerca-de-redsalud/administracion"),
 ("Clínica RedSalud Santiago","Empresas RedSalud","","Av. Libertador Bernardo O'Higgins 4850","Estación Central","+56 600 718 6000","redsalud.cl",
  "","","","","redsalud.trabajando.cl","","","","redsalud.cl"),
 ("Clínica RedSalud Vitacura (ex Tabancura)","Empresas RedSalud","78.053.560-1","Av. Tabancura 1185","Vitacura","+56 2 2395 4000","redsalud.cl",
  "","","","","redsalud.trabajando.cl","","","","superdesalud.gob.cl/registro/clinica-redsalud-vitacura/"),
 ("Hospital Clínico UC Christus","UC Christus Servicios Clínicos SpA","99.573.490-7","Marcoleta 367","Santiago","+56 2 2354 3000","ucchristus.cl",
  "","","","","med.puc.trabajando.cl (formulario: contactenos.ucchristus.cl)","","","","superdesalud.gob.cl/registro/hospital-clinico-uc-red-salud-uc-christus/"),
 ("Clínica UC San Carlos de Apoquindo","UC Christus Servicios Clínicos SpA","99.573.490-7","Camino El Alba 12351","Las Condes","+56 2 2754 8726","ucchristus.cl",
  "","","","","ucsancarlos.trabajando.cl","","","","superdesalud.gob.cl/registro/clinica-san-carlos-de-apoquindo-red-salud-uc-christus/"),
 ("Clínica Colonial","Clínica Colonial S.A.","96.790.040-0","Palacio Riesco 4515","Huechuraba","+56 2 2578 8500","clinicacolonial.cl",
  "contacto@clinicacolonial.cl","","","","clinicacolonial.cl/trabaja-con-nosotros/","","","","clinicacolonial.cl"),
 ("Clínica Hospital del Profesor","Comunidad Hospital del Profesor","53.125.850-9","Av. Libertador Bernardo O'Higgins 4860","Estación Central","+56 2 2299 6300","chp.cl",
  "","","","seleccion@chp.cl","chp.trabajando.cl","","","","chp.cl/web-chp/trabaje-con-nosotros"),
 ("Centros Médicos Vidaintegra","Vidaintegra S.A.","","Varias sedes RM","Santiago","+56 600 600 8432","vidaintegra.cl",
  "","","","postulaciones@vidaintegra.cl","vidaintegra.cl/trabaje-vidaintegra.asp","","","","vidaintegra.cl/contacto.asp"),
 ("Clínica Las Condes","Clínica Las Condes S.A.","93.930.000-7","Estoril 450","Las Condes","+56 2 2610 4000","clinicalascondes.cl",
  "","","","","clinicalascondes.buk.cl/trabaja-con-nosotros","","","","clinicalascondes.cl ; superdesalud.gob.cl/registro/clinica-las-condes/"),
 ("Centro Oftalmológico Láser","Servicios Oftalmológicos CEOLA S.A.","79.859.890-2","Asturias 349","Las Condes","+56 2 3201 2000","centrooftalmologicolaser.cl",
  "ayuda@centrolaser.cl","","","","","","","","centrooftalmologicolaser.cl/contacto/"),
 ("Clínica Oftalmológica Pasteur","Servicios Médicos Luis Pasteur S.A.","78.730.160-6","Av. Luis Pasteur 5917","Vitacura","+56 2 2520 5900","pasteur.cl",
  "contacto@pasteur.cl","","pasteur.cl/ecosistema-pasteur/","","pasteur.cl/trabaja-con-nosotros/","","","","pasteur.cl/contacto/"),
 ("Instituto Oftalmológico Puerta del Sol","Instituto Oftalmológico Puerta del Sol S.A.","96.759.910-7","Puerta del Sol 36","Las Condes","+56 2 2411 5700","puertadelsol.cl",
  "contactoiops@puertadelsol.cl","","","","puertadelsol.cl/trabaja-con-nosotros/","","","","puertadelsol.cl"),
 ("Instituto de Radiomedicina IRAM","Instituto de Radiomedicina Ltda.","85.493.600-K","Av. Américo Vespucio 1314","Vitacura","+56 2 2754 1700","iram.cl",
  "contacto@iram.cl","","","","","","","","iram.cl"),
 ("Hospital Parroquial de San Bernardo","Hospital Parroquial de San Bernardo","82.031.800-5","Av. O'Higgins 04","San Bernardo","+56 2 2373 6500","hpsb.cl",
  "","","hpsb.cl/proveedores/","","hpsb.cl/trabaje-con-nosotros/","","","","hpsb.cl/contacto/"),
 ("Clínica Ensenada","Clínica Ensenada SpA","76.363.205-9","Av. Fermín Vivaceta 957","Independencia","+56 2 2437 3560","clinicaensenada.cl",
  "","","","","Formulario en clinicaensenada.cl","","","","clinicaensenada.cl"),
 ("Nueva Clínica Cordillera","Nueva Clínica Cordillera Prestaciones Hospitalizadas S.A.","76.871.990-K","Av. Alejandro Fleming 7889","Las Condes","+56 2 2834 7500","nuevaclinicacordillera.cl",
  "info@nuevaclinicacordillera.cl","","","","laborum.cl (Red Interclínica)","","","","nuevaclinicacordillera.cl"),
]

HEAD = ["Clínica","Razón social","RUT","Dirección","Comuna","Teléfono","Sitio web","Email general",
        "Email compras/proveedores","Canal compras (portal/nota)","Email RRHH/selección","Canal RRHH/Capacitación",
        "Docencia/Capacitación interna","Nombre contacto","Cargo contacto","Fuente",
        "Nombre decisor (completar)","Cargo decisor (completar)","Email decisor (completar)","Estado seguimiento"]

WIX = ["First Name","Last Name","Email","Phone","Company","Position","Address","Labels"]

def wix_rows():
    for c in C:
        (clin,rs,rut,dir_,com,tel,web,eg,ec,cc,er,cr,doc,nom,cargo,src) = c
        addr = f"{dir_}, {com}, Santiago, Chile"
        base = dict(Company=clin, Phone=tel, Address=addr)
        rows = []
        if ec: rows.append({**base,"Email":ec,"Position":"Compras / Proveedores","Labels":"Clínicas Santiago; Compras"})
        if er: rows.append({**base,"Email":er,"Position":"RRHH / Selección","Labels":"Clínicas Santiago; RRHH"})
        if eg: rows.append({**base,"Email":eg,"Position":"Contacto general","Labels":"Clínicas Santiago; General"})
        if nom:
            parts = nom.split(" ",1)
            rows.append({**base,"First Name":parts[0],"Last Name":parts[1] if len(parts)>1 else "",
                         "Email":"","Position":cargo,"Labels":"Clínicas Santiago; Decisor"})
        if not rows:
            rows.append({**base,"Email":"","Position":"Central telefónica","Labels":"Clínicas Santiago; Sin email"})
        for r in rows:
            yield [r.get(k,"") for k in WIX]

wb = Workbook()
hdr_font = Font(bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="1F4E78")
todo = PatternFill("solid", fgColor="FFF2CC")

def sheet(ws, head, rows, widths=None):
    ws.append(head)
    for cell in ws[1]:
        cell.font, cell.fill = hdr_font, fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows: ws.append(list(r))
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(head, 1):
        col = [str(ws.cell(row=j, column=i).value or "") for j in range(1, ws.max_row+1)]
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(x) for x in col)+2, 50)

ws = wb.active; ws.title = "Clínicas"
sheet(ws, HEAD, [list(c)+["","","","Pendiente"] for c in C])
for row in ws.iter_rows(min_row=2, min_col=17, max_col=19):
    for cell in row: cell.fill = todo

wix = list(wix_rows())
sheet(wb.create_sheet("Wix_Import"), WIX, wix)

notas = wb.create_sheet("Notas")
for line in [
 ["Generado 29-09-2026. Solo datos institucionales publicados en sitios oficiales / Superintendencia de Salud."],
 ["No se incluyeron correos personales de individuos ni correos deducidos por patrón (Ley 19.628 / Ley 21.719)."],
 ["Columnas amarillas en 'Clínicas': completar manualmente (LinkedIn, llamada a central) con el jefe de Compras / Capacitación / Personas."],
 ["Wix: Contactos > Importar > CSV. Usar el archivo clinicas_santiago_wix.csv (UTF-8). Mapear columnas si Wix lo solicita."],
 ["Descartadas: Clínica Las Lilas (aparente quiebra), Santa Sofía (cerrada), Sierra Bella (trasladada). Clínica Las Condes en crisis financiera 2025-2026."],
 ["Verificar antes de usar: seleccion@chp.cl y info@nuevaclinicacordillera.cl (vistos solo en resultados de búsqueda)."],
]: notas.append(line)
notas.column_dimensions["A"].width = 130

wb.save(OUT/"clinicas_santiago_contactos.xlsx")
with open(OUT/"clinicas_santiago_wix.csv","w",newline="",encoding="utf-8-sig") as f:
    w = csv.writer(f); w.writerow(WIX); w.writerows(wix)
print(len(C), "clínicas;", len(wix), "filas Wix")
